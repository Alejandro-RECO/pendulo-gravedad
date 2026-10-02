"""Exporta los datos y el cálculo a Excel, con fórmulas, para revisarlo a mano.

El Excel se genera desde el CSV crudo (la fuente de verdad); no se edita para alimentar el programa.
Uso desde consola:  python exportar_excel.py  -> datos/2026-10-01_pendulo.xlsx
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

RAIZ = Path(__file__).resolve().parent


def _encabezado(ws, columnas, anchos=None):
    ws.append(columnas)
    for i, celda in enumerate(ws[1], start=1):
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E79")
        celda.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[celda.column_letter].width = (anchos or {}).get(i, 15)
    ws.freeze_panes = "A2"


def exportar(r: dict, ruta: Path) -> None:
    m = r["montaje"]
    n = m["n_oscilaciones"]
    wb = Workbook()

    # --- series: una fila por medición
    ws = wb.active
    ws.title = "series"
    _encabezado(ws, ["L (m)", "Medición", "n", "t₁₀ (s)", "T = t₁₀/n (s)", "θ0 (°)", "m (g)", "Uso"],
                {8: 14})
    excluidas = r["series_excluidas"]
    filas_de = {}
    for s in sorted(r["series"] + excluidas, key=lambda s: (-s.L, s.repeticion)):
        f = ws.max_row + 1
        ws.append([s.L, s.repeticion, n, s.t_n(n), f"=D{f}/C{f}", m["amplitud_grados"],
                   m["masa_kg"] * 1000, "validación" if s in excluidas else "ajuste"])
        ws[f"E{f}"].number_format = "0.0000"
        if s not in excluidas:
            filas_de.setdefault(s.L, []).append(f)

    # --- promedios de las longitudes del ajuste, con fórmulas sobre la hoja series
    wp = wb.create_sheet("promedios")
    _encabezado(wp, ["L (m)", "T₁ (s)", "T₂ (s)", "T₃ (s)", "T promedio (s)", "desv. (s)", "T² (s²)",
                     "g = 4π²L/T² (m/s²)"])
    for L in sorted(filas_de, reverse=True):
        f = wp.max_row + 1
        refs = [f"=series!E{x}" for x in filas_de[L]]
        wp.append([L, *refs, f"=AVERAGE(B{f}:D{f})", f"=STDEV(B{f}:D{f})", f"=E{f}^2",
                   f"=4*PI()^2*A{f}/G{f}"])
        for col in "BCDEFG":
            wp[f"{col}{f}"].number_format = "0.0000"
        wp[f"H{f}"].number_format = "0.000"
    ultima = wp.max_row

    # --- ajuste con funciones de Excel (cálculo independiente del programa)
    wa = wb.create_sheet("ajuste")
    _encabezado(wa, ["Magnitud", "Excel", "Programa"], {1: 34, 2: 14, 3: 14})
    res = r["resultado"]
    x, y = f"promedios!A2:A{ultima}", f"promedios!G2:G{ultima}"
    filas = [
        ("Pendiente m de T² vs L (s²/m)", f"=SLOPE({y},{x})", res["m"]),
        ("Corte b (s²)", f"=INTERCEPT({y},{x})", res["b"]),
        ("R²", f"=RSQ({y},{x})", res["r2"]),
        ("g = 4π²/m (m/s²)", "=4*PI()^2/B2", res["g"]),
        ("u(g) del programa (m/s²)", None, res["u_g"]),
        ("g de Bogotá (m/s²)", m["g_referencia"], m["g_referencia"]),
        ("Error % frente a Bogotá", "=ABS(B7-B5)/B7*100", res["error_pct"]),
        ("g de clase (m/s²)", m["g_clase"], m["g_clase"]),
        ("Error % frente a clase", "=ABS(B9-B5)/B9*100", abs(m["g_clase"] - res["g"]) / m["g_clase"] * 100),
    ]
    for nombre, excel, prog in filas:
        wa.append([nombre, excel, prog])
        for col in "BC":
            wa[f"{col}{wa.max_row}"].number_format = "0.0000"

    # --- vueltas: los datos crudos
    wv = wb.create_sheet("vueltas")
    _encabezado(wv, ["L (m)", "Medición", "Vuelta", "Duración de la vuelta (s)", "Tiempo acumulado (s)"])
    for s in sorted(r["series"] + excluidas, key=lambda s: (-s.L, s.repeticion)):
        for k, (v, a) in enumerate(zip(s.vueltas, s.acumulados), start=1):
            wv.append([s.L, s.repeticion, k, v, a])

    # --- instrumentos
    wi = wb.create_sheet("instrumentos")
    _encabezado(wi, ["Instrumento", "Magnitud", "Rango", "Resolución", "Del instrumento", "Efectiva",
                     "Justificación"], {1: 40, 7: 60})
    for i in m["instrumentos"]:
        wi.append([i["instrumento"], i["magnitud"], i["rango"], i["resolucion"], i["u_instrumento"],
                   i["u_efectiva"], i["justificacion"]])
    wi.append([])
    wi.append(["u_g/g = √[(u_L/L)² + (2u_T/T)²]"])

    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


if __name__ == "__main__":
    import sys

    sys.path.insert(0, str(RAIZ))
    from pendulo.analisis import procesar
    from pendulo.datos import cargar_montaje

    destino = RAIZ / "datos" / "2026-10-01_pendulo.xlsx"
    exportar(procesar(cargar_montaje()), destino)
    print(f"Escrito: {destino}")
